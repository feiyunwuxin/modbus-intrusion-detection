"""
generate_proposal_docx.py — 从 proposal.md 生成开题报告 .docx
================================================================
- 4 色诊断框（pitfall / todo / good / grade）
- 一级标题黑体二号 / 二级黑体三号 / 正文宋体小四
- 与 thesis-template.docx 风格统一
- 自动检测 P-P.x 编号

用法：python generate_proposal_docx.py
输出：proposal-template.docx （与本脚本同目录）
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap
import re
import os
import sys

# ---------- 内联样式函数（与 word-template/generate_template.py 同步）----------

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge in kwargs:
            tag = qn(f'w:{edge}')
            elem = tcBorders.find(tag)
            if elem is None:
                elem = OxmlElement(f'w:{edge}')
                tcBorders.append(elem)
            for k, v in kwargs[edge].items():
                elem.set(qn(f'w:{k}'), v)


def set_cell_shading(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def add_diagnostic_box(doc, kind, code, text):
    config = {
        'pitfall': {'border': 'C00000', 'shade': 'FFEBEB', 'icon': '[X]', 'tag': f'[pitfall {code}]', 'font': '楷体'},
        'todo':    {'border': 'BF8F00', 'shade': 'FFFACD', 'icon': '[?]', 'tag': f'[todo {code}]',    'font': '楷体'},
        'good':    {'border': '00B050', 'shade': 'E2EFDA', 'icon': '[OK]', 'tag': f'[good {code}]',   'font': '楷体'},
        'grade':   {'border': '2E75B6', 'shade': 'EBF5FF', 'icon': '[T]', 'tag': f'[grade {code}]',  'font': '楷体'},
    }[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_border(cell,
                    top={'val': 'single', 'sz': '12', 'color': config['border']},
                    left={'val': 'single', 'sz': '12', 'color': config['border']},
                    bottom={'val': 'single', 'sz': '12', 'color': config['border']},
                    right={'val': 'single', 'sz': '12', 'color': config['border']})
    set_cell_shading(cell, config['shade'])
    p = cell.paragraphs[0]
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(f"{config['tag']}  {text}")
    run.font.size = Pt(10.5)
    run.font.name = config['font']
    doc.add_paragraph()


def add_styled_paragraph(doc, text, style='Normal', align=None, indent=True):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if indent and style == 'Normal':
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    if style == 'Normal':
        run.font.size = Pt(12)
        run.font.name = '宋体'
    return p


def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.name = '黑体'
    return p


def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.name = '黑体'
    return p


def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.name = '黑体'
    return p


def setup_page(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)
    section.gutter = Cm(0.5)

# ---------- 主生成 ----------
doc = Document()
setup_page(doc)

# ========== 封面 ==========
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
run = p.add_run('（学校名）')
run.font.size = Pt(16)
run.font.name = '黑体'

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
run = p.add_run('硕士研究生开题报告')
run.font.size = Pt(36)
run.font.bold = True
run.font.name = '黑体'

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(48)
run = p.add_run('题目：基于 TCN-SE 与混合精度量化的\nSCADA 入侵检测研究')
run.font.size = Pt(22)
run.font.name = '黑体'

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(120)
run = p.add_run('研究生姓名：（姓名）\n指导教师：（教授）\n专业：（控制工程）\n研究方向：工业控制系统安全\n日期：2026 年 9 月')
run.font.size = Pt(14)
run.font.name = '宋体'

doc.add_page_break()

# ========== 解析 proposal.md 并渲染 ==========
md_path = os.path.join(os.path.dirname(__file__), 'proposal.md')
md_text = open(md_path, encoding='utf-8').read()

# 切分 md 为 block（按空行）
blocks = re.split(r'\n\s*\n', md_text)

DIAG_PATTERN = re.compile(
    r'>\s*(⚠|📝|✓|✏️)\s*\*\*(【(?:学生常见错误|待补充[（(]学生[)）]?|优秀示例|开题错误|开题总评|老师评语)[^】]*】)\*\*\s*(.*?)(?=\n|$)',
    re.DOTALL
)

for block in blocks:
    block = block.strip()
    if not block:
        continue
    if block.startswith('# '):
        # 跳过 H1（封面已处理）
        continue
    if block.startswith('## '):
        text = block[3:].strip()
        add_heading_1(doc, text)
    elif block.startswith('### '):
        text = block[4:].strip()
        add_heading_2(doc, text)
    elif block.startswith('> '):
        # 诊断框
        m = DIAG_PATTERN.search(block)
        if m:
            icon, tag, desc = m.group(1), m.group(2), m.group(3).strip()
            # 判断类型
            if '学生常见错误' in tag or '开题错误' in tag:
                kind = 'pitfall'
            elif '待补充' in tag:
                kind = 'todo'
            elif '优秀示例' in tag:
                kind = 'good'
            elif '评语' in tag or '总评' in tag:
                kind = 'grade'
            else:
                kind = 'pitfall'
            # 提取 P-P.x 编号
            code_match = re.search(r'P-[\d.P]+', tag)
            code = code_match.group(0) if code_match else ''
            add_diagnostic_box(doc, kind, code, desc)
        else:
            # 普通引用块
            for line in block.split('\n'):
                add_styled_paragraph(doc, line.lstrip('> ').strip(), indent=False)
    elif block.startswith('|'):
        # 表格
        rows = [r.strip() for r in block.split('\n') if r.strip().startswith('|')]
        if len(rows) < 2:
            continue
        # 解析表格
        cells_list = []
        for r in rows:
            cells = [c.strip() for c in r.strip('|').split('|')]
            cells_list.append(cells)
        # 第二行是分隔符（---|---|---），跳过
        if len(cells_list) >= 2 and re.match(r'^[\s\-:|]+$', cells_list[1][0]):
            data_rows = [cells_list[0]] + cells_list[2:]
        else:
            data_rows = cells_list
        n_rows = len(data_rows)
        n_cols = max(len(r) for r in data_rows)
        table = doc.add_table(rows=n_rows, cols=n_cols)
        table.style = 'Table Grid'
        for i, r in enumerate(data_rows):
            for j in range(n_cols):
                cell = table.rows[i].cells[j] if j < len(r) else table.rows[i].cells[j]
                cell.text = r[j] if j < len(r) else ''
    elif block.startswith('```'):
        # 代码块
        lines = block.strip('`').split('\n')
        code_text = '\n'.join(lines[1:]) if len(lines) > 1 else lines[0]
        p = doc.add_paragraph()
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
    elif block.startswith('---'):
        doc.add_page_break()
    elif re.match(r'^\d+\.\s', block) or re.match(r'^- ', block):
        # 列表
        for line in block.split('\n'):
            line = line.strip()
            if line:
                add_styled_paragraph(doc, line, indent=False)
    else:
        # 普通段落（可能多行）
        for line in block.split('\n'):
            line = line.strip()
            if line:
                add_styled_paragraph(doc, line)

# ---------- 保存 ----------
out_path = os.path.join(os.path.dirname(__file__), 'proposal-template.docx')
doc.save(out_path)
print(f'OK: {out_path}')
print(f'  - Size: {os.path.getsize(out_path) / 1024:.1f} KB')
print(f'  - Paragraphs: {len(doc.paragraphs)}')
print(f'  - Tables: {len(doc.tables)}')
